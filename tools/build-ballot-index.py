#!/usr/bin/env python3
"""Builds the ballot index and the story backfill from the newsroom's spreadsheet.

    python3 tools/build-ballot-index.py [path/to/Measure_Candidate_Backfill.xlsx]

Writes
  election-dashboard/data/ballot-index.json   every measure, contested race and uncontested race, per county

The matcher uses ballot-index.json to decide which measure or race a newly published story is about.
(The spreadsheet's Story columns are read into tools/story-backfill.json for reference only.)
"""
import json, os, re, sys, unicodedata
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'tools/source-data/Measure_Candidate_Backfill.xlsx')
OUT = os.path.join(ROOT, 'election-dashboard/data')

COUNTY_SLUGS = {
    'alameda county': 'alameda', 'contra costa county': 'contra-costa', 'marin county': 'marin',
    'mendocino county': 'mendocino', 'monterey county': 'monterey', 'napa county': 'napa',
    'city and county of san francisco': 'san-francisco', 'san francisco county': 'san-francisco', 'san francisco': 'san-francisco',
    'san joaquin county': 'san-joaquin', 'san mateo county': 'san-mateo', 'santa clara county': 'santa-clara',
    'santa cruz county': 'santa-cruz', 'solano county': 'solano', 'sonoma county': 'sonoma',
}

def clean(s):
    if s is None: return ''
    s = str(s)
    # the sheet has a few mojibake sequences from a Windows export
    s = s.replace('Ã¢??', '—').replace('Ã©', 'é').replace('Ã±', 'ñ').replace('Ã³', 'ó').replace('Ã¡', 'á')
    return re.sub(r'\s+', ' ', s).strip()

def norm(s):
    s = unicodedata.normalize('NFKD', clean(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9]+', ' ', s.lower())).strip()

def rows(ws):
    hdr = [clean(c) for c in next(ws.iter_rows(values_only=True))]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(v for v in r): continue
        yield {hdr[i]: clean(v) for i, v in enumerate(r) if i < len(hdr) and hdr[i]}

def story_urls(row):
    out = []
    for k, v in row.items():
        if k.startswith('Story') and v and v.startswith('http'):
            out.append(v.strip())
    return out

wb = openpyxl.load_workbook(SRC, data_only=True)
items = {}       # id -> item
stories = {}     # url -> set(ids)
def add(item, urls):
    items[item['id']] = item
    for u in urls:
        stories.setdefault(u, set()).add(item['id'])

# Measures: one item per measure
for r in rows(wb['Measures']):
    county = COUNTY_SLUGS[r['County'].lower()]
    name = r['Measure Name']
    m = re.match(r'(Measure|Prop(?:osition)?\.?)\s+([A-Z]{1,3}|\d{1,3})\b', name, re.I)
    letter = m.group(2).upper() if m else ''
    item = {'id': 'm:%s:%s' % (county, norm(name)), 'county': county, 'type': 'measure', 'key': name, 'keyNorm': norm(name),
            'letter': letter, 'jurisdiction': r.get('Measure Juristiction', ''), 'desc': r.get('Measure Description', '')[:400]}
    add(item, story_urls(r))

# Contested races: one item per race, with its candidates
races = {}
for r in rows(wb['Candidates']):
    county = COUNTY_SLUGS[r['County'].lower()]
    rid = 'r:%s:%s' % (county, norm(r['Race']))
    it = races.setdefault(rid, {'id': rid, 'county': county, 'type': 'race', 'key': r['Race'], 'keyNorm': norm(r['Race']), 'candidates': []})
    it['candidates'].append(r['Candidate Name'])
    stories.setdefault
    for u in story_urls(r):
        stories.setdefault(u, set()).add(rid)
for it in races.values(): items[it['id']] = it

# Uncontested: one item per race as well
unc = {}
for r in rows(wb['Uncontested Races']):
    county = COUNTY_SLUGS[r['County'].lower()]
    rid = 'u:%s:%s' % (county, norm(r['Race']))
    it = unc.setdefault(rid, {'id': rid, 'county': county, 'type': 'uncontested', 'key': r['Race'], 'keyNorm': norm(r['Race']), 'candidates': []})
    it['candidates'].append(r['Candidate Name'])
    for u in story_urls(r):
        stories.setdefault(u, set()).add(rid)
for it in unc.values(): items[it['id']] = it

index = {'generated_from': os.path.basename(SRC), 'counties': sorted(set(COUNTY_SLUGS.values())), 'items': sorted(items.values(), key=lambda i: (i['county'], i['type'], i['key']))}
os.makedirs(OUT, exist_ok=True)
json.dump(index, open(os.path.join(OUT, 'ballot-index.json'), 'w'), ensure_ascii=False, indent=0)
json.dump({u: {'items': sorted(ids)} for u, ids in sorted(stories.items())}, open(os.path.join(ROOT, 'tools/story-backfill.json'), 'w'), ensure_ascii=False, indent=1)
by_type = {}
for i in items.values(): by_type[i['type']] = by_type.get(i['type'], 0) + 1
print('items:', by_type, '| stories:', len(stories), '| index bytes:', os.path.getsize(os.path.join(OUT, 'ballot-index.json')))
